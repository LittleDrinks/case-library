"""Word 编号来源与原生 REF 交叉引用。"""

from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.modules.documents.styles import set_run_font


def _bookmark_name(number: int) -> str:
    return f"_RefCaseSource{number}"


def _field_boundary(paragraph: Paragraph, kind: str) -> None:
    boundary = OxmlElement("w:fldChar")
    boundary.set(qn("w:fldCharType"), kind)
    if kind == "begin":
        boundary.set(qn("w:dirty"), "true")
    paragraph.add_run()._r.append(boundary)


def add_citation(paragraph: Paragraph, number: int, font: str, size: int) -> None:
    _field_boundary(paragraph, "begin")
    code = paragraph.add_run()
    set_run_font(code, font, size)
    code.font.superscript = True
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    # CHARFORMAT 取 REF 首字母的格式，更新域后仍保留上标。
    instruction.text = f" REF {_bookmark_name(number)} \\n \\h \\* CHARFORMAT "
    code._r.append(instruction)
    _field_boundary(paragraph, "separate")
    result = paragraph.add_run(f"〔{number}〕")
    set_run_font(result, font, size)
    result.font.superscript = True
    _field_boundary(paragraph, "end")


def bookmark_reference(paragraph: Paragraph, number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), _bookmark_name(number))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.append(end)


def configure_reference_numbering(document: DocxDocument, num_id: int) -> None:
    numbering = document.part.numbering_part.element
    abstract_id = numbering.num_having_numId(num_id).abstractNumId.val
    abstract = next(item for item in numbering.findall(qn("w:abstractNum"))
                    if int(item.get(qn("w:abstractNumId"))) == abstract_id)
    level = abstract.find(qn("w:lvl"))
    level.find(qn("w:lvlText")).set(qn("w:val"), "〔%1〕")
    # 清除列表样式关联，编号只属于本次导出的来源清单。
    style = level.find(qn("w:pStyle"))
    if style is not None:
        level.remove(style)
    properties = level.find(qn("w:pPr"))
    indent = properties.find(qn("w:ind"))
    indent.set(qn("w:left"), "600")
    indent.set(qn("w:hanging"), "600")
    properties.find(f"{qn('w:tabs')}/{qn('w:tab')}").set(qn("w:pos"), "600")
